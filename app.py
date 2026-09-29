import streamlit as st
import pandas as pd
import datetime
from pptx import Presentation
from pptx.util import Inches, Pt
import io

# =====================================================================
# 1. CORE SYSTEM SETUP & BRANDING
# =====================================================================
st.set_page_config(page_title="Advanced End-to-End Manufacturing ERP", layout="wide", page_icon="⚙️")

# Master Safety Gate Password
MASTER_PASSWORD = "admin_erp_2026"

def check_password():
    if "authenticated" not in st.session_state:
        st.session_state.authenticated = False
    if not st.session_state.authenticated:
        st.title("🔒 Enterprise Manufacturing Execution Portal (MES)")
        user_pass = st.text_input("Enter Private ERP Master Security Key:", type="password")
        if st.button("Unlock Enterprise Systems"):
            if user_pass == MASTER_PASSWORD:
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("Invalid Security Key. Access Revoked.")
        return False
    return True

if check_password():
    # =====================================================================
    # 2. STATE DATA REGISTRIES (Simulated Central Cloud Relational Database)
    # =====================================================================
    # RM/PM Inventory Ledger
    if "inventory_db" not in st.session_state:
        st.session_state.inventory_db = pd.DataFrame([
            {"Material Code": "RM-STEEL-01", "Material Name": "Premium Structural Steel", "Type": "Raw Material (RM)", "Stock Qty": 1000, "Unit": "Kgs"},
            {"Material Code": "PM-BOX-05", "Material Name": "Industrial Storage Pallet Boxes", "Type": "Packing Material (PM)", "Stock Qty": 150, "Unit": "Pcs"},
            {"Material Code": "RM-ALUM-02", "Material Name": "Tempered Aluminum Sheeting", "Type": "Raw Material (RM)", "Stock Qty": 500, "Unit": "Sheets"}
        ])

    # Purchase Orders (PO) Registry
    if "po_db" not in st.session_state:
        st.session_state.po_db = pd.DataFrame([
            {"PO Number": "PO-2026-001", "Vendor": "Apex Metal Forges", "Material Code": "RM-STEEL-01", "Ordered Qty": 500, "Status": "Open"}
        ])

    # GRN Logs
    if "grn_db" not in st.session_state:
        st.session_state.grn_db = pd.DataFrame([
            {"GRN Number": "GRN-9901", "PO Reference": "PO-2026-001", "Date": "2026-09-20", "Received Qty": 500, "Inspected Status": "Accepted"}
        ])

    # Bill of Materials (BOM) Recipes
    if "bom_db" not in st.session_state:
        st.session_state.bom_db = pd.DataFrame([
            {"Finished Good": "Heavy Duty Steel Chassis", "Component Code": "RM-STEEL-01", "Qty Required": 5, "Component Unit": "Kgs"},
            {"Finished Good": "Heavy Duty Steel Chassis", "Component Code": "PM-BOX-05", "Qty Required": 1, "Component Unit": "Pcs"}
        ])

    # Production Batches
    if "batch_db" not in st.session_state:
        st.session_state.batch_db = pd.DataFrame([
            {"Batch ID": "B-BATCH-101", "Target Product": "Heavy Duty Steel Chassis", "Batch Target Qty": 20, "Status": "Issued & Released"}
        ])

    # Quality Assurance & Output Yield logs
    if "production_output_db" not in st.session_state:
        st.session_state.production_output_db = pd.DataFrame([
            {"Batch ID": "B-BATCH-101", "Produced Qty": 18, "Rejected Qty": 2, "Calculated Yield": "90.0%"}
        ])

    # Sales Invoices & Dispatches Registry
    if "sales_db" not in st.session_state:
        st.session_state.sales_db = pd.DataFrame([
            {"Invoice ID": "INV-5501", "Client Account": "Titan Logistics Corp", "Product Dispatched": "Heavy Duty Steel Chassis", "Shipped Qty": 10, "Revenue Realized ($)": 4500.00}
        ])

    # Finished Goods Inventory tracking
    if "fg_inventory" not in st.session_state:
        st.session_state.fg_inventory = {"Heavy Duty Steel Chassis": 8}

    # =====================================================================
    # 3. INTERACTIVE NAVIGATION SIDEBAR
    # =====================================================================
    st.sidebar.title("🏭 Factory Operations Management")
    module = st.sidebar.radio("Select Processing Phase:", [
        "📊 Master Executive Overview",
        "1️⃣ Procurement & Purchase Orders",
        "2️⃣ Inward Logistics & GRN",
        "3️⃣ Master BOM Configuration",
        "4️⃣ Batch Issuance & Material Release",
        "5️⃣ Production Output & Yield Engine",
        "6️⃣ Dispatches & Sales Invoices"
    ])

    # =====================================================================
    # MODULE: EXECUTIVE OVERVIEW & MASTER REPORTING
    # =====================================================================
    if module == "📊 Master Executive Overview":
        st.title("📊 Master Executive Overview Dashboard")
        st.write(f"System State Synchronized Live on: **{datetime.date.today()}**")
        st.markdown("---")
        
        # Display operational status tickers
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("Open Purchase Orders", f"{len(st.session_state.po_db[st.session_state.po_db['Status'] == 'Open'])} Active")
        k2.metric("Active Work Batches", f"{len(st.session_state.batch_db[st.session_state.batch_db['Status'] == 'Issued & Released'])} Running")
        
        all_sales_revenue = st.session_state.sales_db["Revenue Realized ($)"].sum()
        k3.metric("Gross Invoiced Sales", f"${all_sales_revenue:,.2f}")
        
        total_p = st.session_state.production_output_db["Produced Qty"].sum()
        total_r = st.session_state.production_output_db["Rejected Qty"].sum()
        plant_yield = (total_p / (total_p + total_r) * 100) if (total_p + total_r) > 0 else 100.0
        k4.metric("Average Shop Floor Yield", f"{plant_yield:.1f}%")

        st.markdown("---")
        col_left, col_right = st.columns(2)
        with col_left:
            st.subheader("📦 Live RM/PM Stocks")
            st.dataframe(st.session_state.inventory_db, use_container_width=True, hide_index=True)
        with col_right:
            st.subheader("🏆 Vaulted Finished Goods Available for Sale")
            for product, qty in st.session_state.fg_inventory.items():
                st.info(f"**{product}**: {qty} Units available in Finished Goods Warehouse.")

        st.markdown("---")
        st.subheader("💼 Compile Presentations")
        if st.button("🚀 Auto-Generate PowerPoint Operational Report"):
            prs = Presentation()
            slide_layout = prs.slide_layouts
            
            # Slide 1: Main Title
            s1 = prs.slides.add_slide(slide_layout)
            s1.shapes.title.text = "Operational Manufacturing Execution Report"
            s1.placeholders.text = f"Data Analytics Period: {datetime.date.today()}"
            
            # Slide 2: Data Metrics Text Box
            s2 = prs.slides.add_slide(slide_layout)
            s2.shapes.title.text = "Key Operational Throughput Metrics"
            tb = s2.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(9), Inches(4)).text_frame
            tb.word_wrap = True
            tb.paragraphs.text = f"• Total Factory Yield Rate Calculated: {plant_yield:.2f}%"
            tb.add_paragraph().text = f"• Gross Financial Revenue Invoiced: ${all_sales_revenue:,.2f}"
            tb.add_paragraph().text = f"• Total System Logged Purchase Orders: {len(st.session_state.po_db)} Registered"
            
            # Slide 3: Tabular Breakdown of Stock
            s3 = prs.slides.add_slide(slide_layout)
            s3.shapes.title.text = "Current Raw Supply Balances"
            df_e = st.session_state.inventory_db
            r, c = len(df_e) + 1, len(df_e.columns)
            table_shape = s3.shapes.add_table(r, c, Inches(0.5), Inches(1.8), Inches(9), Inches(3))
            tbl = table_shape.table
            for j, col_n in enumerate(df_e.columns):
                tbl.cell(0, j).text = str(col_n)
            for i, r_data in df_e.iterrows():
                for j, val in enumerate(r_data):
                    tbl.cell(i+1, j).text = str(val)
            
            buf = io.BytesIO()
            prs.save(buf)
            buf.seek(0)
            st.download_button(label="📥 Download Executive Presentation", data=buf, file_name="Operations_Executive_Review.pptx")

    # =====================================================================
    # 1. PROCUREMENT & PURCHASE ORDERS
    # =====================================================================
    elif module == "1️⃣ Procurement & Purchase Orders":
        st.title("1️⃣ Purchase Order (PO) Engine")
        st.write("Authorize external purchase commitments for RM/PM materials to registered vendors.")
        
        c1, c2 = st.columns([1, 2])
        with c1:
            st.subheader("Generate Fresh PO")
            with st.form("po_form", clear_on_submit=True):
                po_id = f"PO-2026-{len(st.session_state.po_db) + 1:03d}"
                st.warning(f"Drafting Reference Number: {po_id}")
                v_name = st.text_input("Target Vendor/Supplier Name", placeholder="e.g., Nippon Steel Corp")
                m_code = st.selectbox("Select Target Material Code", st.session_state.inventory_db["Material Code"].tolist())
                o_qty = st.number_input("Purchase Commitment Quantity", min_value=1, value=100)
                
                if st.form_submit_button("Authorize & Release Purchase Order"):
                    if v_name:
                        new_p = {"PO Number": po_id, "Vendor": v_name, "Material Code": m_code, "Ordered Qty": o_qty, "Status": "Open"}
                        st.session_state.po_db = pd.concat([st.session_state.po_db, pd.DataFrame([new_p])], ignore_index=True)
                        st.success(f"PO {po_id} has been formally released to {v_name}!")
                        st.rerun()
                    else:
